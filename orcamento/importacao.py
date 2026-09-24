"""Importação de orçamento a partir de planilha Excel (.xlsx).

A planilha pode ter qualquer layout, desde que tenha uma linha de cabeçalho
com, pelo menos, as colunas de descrição, quantidade e preço unitário. As
colunas são reconhecidas pelo nome (ex.: "Item", "Descrição", "Und",
"Quant.", "Preço unitário"). A EAP é montada pela numeração da coluna Item:
linhas sem quantidade são etapas (1, 1.1...) e linhas com quantidade são
itens, colocados na etapa de numeração mais próxima (1.1.3 vai para 1.1).
"""

import re
import unicodedata
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction

from cadastros.models import Composicao, Insumo, UnidadeMedida

from .models import Etapa, ItemOrcamento, Orcamento, arredondar

ZERO = Decimal("0")
LINHAS_PARA_ACHAR_CABECALHO = 40
CODIGO_SEM_ETAPA = "00"


def normalizar(texto):
    """'Preço Unit. (R$)' -> 'preco unit r'"""
    texto = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texto.lower()).split())


def identificar_coluna(cabecalho):
    """Qual informação a coluna traz, pelo texto do cabeçalho (ou None)."""
    h = normalizar(cabecalho)
    if not h:
        return None
    palavras = h.split()
    if "unit" in h or "unitario" in palavras:
        return "preco_com_bdi" if "com bdi" in h else "preco"
    if palavras[0] in {"total", "subtotal"} or h.startswith(("valor total", "preco total", "custo total")):
        return "total"
    if h.startswith(("descri", "discrimina", "especifica", "servico")):
        return "descricao"
    if h in {"unidade", "und", "un", "unid", "u", "unidade de medida"}:
        return "unidade"
    if h.startswith("quant") or h in {"qtd", "qtde", "qde", "qt"}:
        return "quantidade"
    if h in {"item", "itens", "eap", "n", "no", "num", "numero", "cod item", "codigo eap", "item eap"}:
        return "codigo"
    if h in {"codigo", "cod", "referencia", "ref", "composicao", "codigo sinapi", "sinapi", "cod sinapi",
             "codigo da composicao", "cod composicao"}:
        return "referencia"
    return None


def formatar(valor, casas=2):
    """1234.5 -> '1.234,50' (para mensagens)."""
    texto = f"{valor:,.{casas}f}"
    return texto.replace(",", "_").replace(".", ",").replace("_", ".")


def ler_numero(valor):
    """Número de uma célula: aceita número do Excel, '1.234,56', 'R$ 10,5' e '1234.56'."""
    if valor is None:
        return None
    if isinstance(valor, bool):
        raise ValueError(valor)
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    texto = str(valor).replace("R$", "").replace(" ", "").strip()
    if not texto or texto in {"-", "–"}:
        return None
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return Decimal(texto.replace(" ", ""))
    except InvalidOperation:
        raise ValueError(valor) from None


def ler_codigo(valor):
    """Numeração da EAP como texto: 1 -> '1'; 1.1 -> '1.1'; ' 01.02. ' -> '01.02'."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip().rstrip(".")


def ler_texto(valor):
    return " ".join(str(valor).split()) if valor is not None else ""


def pai_de(codigo, existentes):
    """Etapa mais próxima acima do código: '1.2.3' procura '1.2' e depois '1'."""
    partes = codigo.split(".")
    for tamanho in range(len(partes) - 1, 0, -1):
        candidato = ".".join(partes[:tamanho])
        if candidato in existentes:
            return candidato
    return None


@dataclass
class EtapaLida:
    codigo: str
    descricao: str
    pai: str | None
    linha: int


@dataclass
class ItemLido:
    etapa: str
    codigo: str
    referencia: str
    descricao: str
    unidade: str
    quantidade: str
    preco: str | None
    linha: int


@dataclass
class Leitura:
    aba: str = ""
    colunas: dict = field(default_factory=dict)
    etapas: list = field(default_factory=list)
    itens: list = field(default_factory=list)
    erros: list = field(default_factory=list)
    avisos: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.erros and bool(self.itens)

    def para_sessao(self):
        return {
            "aba": self.aba,
            "colunas": self.colunas,
            "etapas": [asdict(e) for e in self.etapas],
            "itens": [asdict(i) for i in self.itens],
            "avisos": self.avisos,
        }

    @classmethod
    def da_sessao(cls, dados):
        return cls(
            aba=dados["aba"],
            colunas=dados["colunas"],
            etapas=[EtapaLida(**e) for e in dados["etapas"]],
            itens=[ItemLido(**i) for i in dados["itens"]],
            avisos=dados["avisos"],
        )

    def resumo(self):
        """Etapas em ordem, com nível, quantidade de itens e total (incluindo subetapas)."""
        total_direto, contagem = {}, {}
        for item in self.itens:
            valor = Decimal(item.quantidade) * Decimal(item.preco or "0")
            total_direto[item.etapa] = total_direto.get(item.etapa, ZERO) + valor
            contagem[item.etapa] = contagem.get(item.etapa, 0) + 1
        linhas = []
        for etapa in self.etapas:
            nivel, pai = 0, etapa.pai
            while pai:
                nivel += 1
                pai = next(e.pai for e in self.etapas if e.codigo == pai)
            prefixo = etapa.codigo + "."
            total = sum(
                (v for c, v in total_direto.items() if c == etapa.codigo or c.startswith(prefixo)), ZERO
            )
            linhas.append({
                "etapa": etapa, "nivel": nivel, "itens": contagem.get(etapa.codigo, 0),
                "total": arredondar(total),
            })
        return linhas

    @property
    def total(self):
        return arredondar(sum((Decimal(i.quantidade) * Decimal(i.preco or "0") for i in self.itens), ZERO))


def achar_cabecalho(aba):
    """(número da linha, {campo: índice da coluna}) ou (None, None)."""
    for numero, linha in enumerate(aba.iter_rows(max_row=LINHAS_PARA_ACHAR_CABECALHO, values_only=True), 1):
        colunas = {}
        for indice, celula in enumerate(linha):
            campo = identificar_coluna(celula)
            if campo and campo not in colunas:
                colunas[campo] = indice
        if {"descricao", "quantidade"} <= colunas.keys() and ("preco" in colunas or "preco_com_bdi" in colunas):
            if "preco" not in colunas:
                colunas["preco"] = colunas.pop("preco_com_bdi")
            colunas.pop("preco_com_bdi", None)
            if "codigo" not in colunas and "referencia" in colunas:
                # Só há uma coluna de código: ela é a numeração da EAP.
                colunas["codigo"] = colunas.pop("referencia")
            return numero, colunas
    return None, None


def ler_planilha(arquivo):
    leitura = Leitura()
    try:
        livro = openpyxl.load_workbook(arquivo, data_only=True, read_only=True)
    except Exception:
        leitura.erros.append("Não foi possível abrir o arquivo. Envie uma planilha do Excel no formato .xlsx.")
        return leitura

    for aba in livro.worksheets:
        numero_cabecalho, colunas = achar_cabecalho(aba)
        if colunas:
            break
    else:
        leitura.erros.append(
            "Não encontrei a linha de cabeçalho. A planilha precisa ter colunas de descrição, "
            "quantidade e preço unitário (ex.: \"Descrição\", \"Quant.\", \"Preço unitário\")."
        )
        return leitura

    leitura.aba = aba.title
    nomes = {"codigo": "Item", "referencia": "Código", "descricao": "Descrição", "unidade": "Unidade",
             "quantidade": "Quantidade", "preco": "Preço unitário", "total": "Total"}
    leitura.colunas = {nomes[c]: openpyxl.utils.get_column_letter(i + 1) for c, i in colunas.items() if c in nomes}
    if "codigo" not in colunas:
        leitura.avisos.append(
            "A planilha não tem coluna \"Item\" com a numeração da EAP; as linhas sem quantidade "
            "viram etapas numeradas em sequência e cada item fica na última etapa acima dele."
        )

    etapas = {}
    ultima_etapa = None
    sequencia = 0

    def celula(linha, campo):
        indice = colunas.get(campo)
        return linha[indice] if indice is not None and indice < len(linha) else None

    for numero, linha in enumerate(aba.iter_rows(min_row=numero_cabecalho + 1, values_only=True),
                                   numero_cabecalho + 1):
        codigo = ler_codigo(celula(linha, "codigo"))
        descricao = ler_texto(celula(linha, "descricao"))
        try:
            quantidade = ler_numero(celula(linha, "quantidade"))
            preco = ler_numero(celula(linha, "preco"))
            total = ler_numero(celula(linha, "total"))
        except ValueError as erro:
            leitura.erros.append(f"Linha {numero}: valor numérico inválido ({erro.args[0]!r}).")
            continue

        if not codigo and not descricao and quantidade is None:
            continue  # linha em branco
        if not codigo and re.match(r"^(sub)?total|^bdi\b|^valor total", normalizar(descricao)):
            continue  # linha de totalização

        if quantidade is None:
            # Etapa (título de grupo).
            if not descricao:
                continue
            if not codigo:
                if "codigo" in colunas:
                    leitura.avisos.append(f"Linha {numero} ignorada: sem item e sem quantidade ({descricao[:60]}).")
                    continue
                sequencia += 1
                codigo = f"{sequencia:02d}"
            if codigo in etapas:
                leitura.erros.append(f"Linha {numero}: a etapa {codigo} aparece mais de uma vez.")
                continue
            if len(codigo) > 30:
                leitura.erros.append(f"Linha {numero}: numeração longa demais ({codigo}).")
                continue
            pai = pai_de(codigo, etapas)
            if "." in codigo and pai is None:
                leitura.avisos.append(f"Linha {numero}: a etapa {codigo} não tem etapa acima; ficou no primeiro nível.")
            etapas[codigo] = EtapaLida(codigo, descricao[:255], pai, numero)
            ultima_etapa = codigo
            continue

        # Item.
        if not descricao:
            leitura.erros.append(f"Linha {numero}: item sem descrição.")
            continue
        if quantidade < 0 or (preco is not None and preco < 0):
            leitura.erros.append(f"Linha {numero}: quantidade ou preço negativo.")
            continue
        if quantidade == 0:
            leitura.avisos.append(f"Linha {numero} ignorada: quantidade zero ({descricao[:60]}).")
            continue
        referencia = ler_codigo(celula(linha, "referencia"))
        if preco is None and not referencia:
            leitura.erros.append(f"Linha {numero}: item sem preço unitário ({descricao[:60]}).")
            continue
        etapa = (pai_de(codigo, etapas) if codigo else None) or ultima_etapa
        if etapa is None:
            if CODIGO_SEM_ETAPA not in etapas:
                etapas[CODIGO_SEM_ETAPA] = EtapaLida(CODIGO_SEM_ETAPA, "Itens sem etapa", None, numero)
                leitura.avisos.append("Há itens antes de qualquer etapa; eles foram para a etapa 00 \"Itens sem etapa\".")
            etapa = CODIGO_SEM_ETAPA
        unidade = ler_texto(celula(linha, "unidade"))[:10]
        if total is not None and preco is not None:
            calculado = quantidade * preco
            if abs(calculado - total) > max(Decimal("0.05"), abs(total) * Decimal("0.005")):
                leitura.avisos.append(
                    f"Linha {numero}: quantidade × preço = {formatar(calculado)}, mas o total da planilha é "
                    f"{formatar(total)}. Confira se a coluna de preço é a certa (com ou sem BDI)."
                )
        leitura.itens.append(ItemLido(
            etapa=etapa, codigo=codigo[:30], referencia=referencia[:30], descricao=descricao[:500],
            unidade=unidade, quantidade=str(quantidade), preco=str(preco) if preco is not None else None,
            linha=numero,
        ))

    leitura.etapas = list(etapas.values())
    if not leitura.itens and not leitura.erros:
        leitura.erros.append("Nenhum item com quantidade foi encontrado abaixo do cabeçalho.")
    return leitura


def conferir_cadastros(leitura):
    """Liga os códigos da planilha aos cadastros e aponta itens sem preço nem cadastro."""
    codigos = {c.lower() for c in Composicao.objects.values_list("codigo", flat=True)}
    codigos |= {c.lower() for c in Insumo.objects.values_list("codigo", flat=True)}
    ligados = 0
    for item in leitura.itens:
        encontrado = bool(item.referencia) and item.referencia.lower() in codigos
        ligados += encontrado
        if item.preco is None and not encontrado:
            leitura.erros.append(
                f"Linha {item.linha}: item sem preço e o código {item.referencia} não está cadastrado."
            )
    if ligados:
        leitura.avisos.append(
            f"{ligados} item(ns) ligados a composições ou insumos já cadastrados pelo código."
        )


@transaction.atomic
def importar(leitura, obra, descricao, data_base, bdi_percentual, versao=None):
    """Cria o orçamento com as etapas e os itens lidos. Devolve o orçamento."""
    if versao is None:
        ultima = obra.orcamentos.order_by("-versao").values_list("versao", flat=True).first()
        versao = (ultima or 0) + 1
    orcamento = Orcamento.objects.create(
        obra=obra, descricao=descricao, versao=versao, data_base=data_base, bdi_percentual=bdi_percentual
    )
    etapas = {}
    for lida in leitura.etapas:
        etapas[lida.codigo] = Etapa.objects.create(
            orcamento=orcamento, codigo=lida.codigo, descricao=lida.descricao,
            pai=etapas.get(lida.pai),
        )

    composicoes = {c.codigo.lower(): c for c in Composicao.objects.all()}
    insumos = {i.codigo.lower(): i for i in Insumo.objects.all()}
    unidades = {u.sigla.lower(): u for u in UnidadeMedida.objects.all()}

    def unidade(sigla):
        sigla = sigla or "un"
        if sigla.lower() not in unidades:
            unidades[sigla.lower()] = UnidadeMedida.objects.create(sigla=sigla, descricao=sigla)
        return unidades[sigla.lower()]

    for lido in leitura.itens:
        ref = lido.referencia.lower()
        composicao = composicoes.get(ref) if ref else None
        insumo = insumos.get(ref) if ref and composicao is None else None
        item = ItemOrcamento(
            etapa=etapas[lido.etapa],
            composicao=composicao,
            insumo=insumo,
            codigo=lido.codigo,
            descricao=lido.descricao,
            unidade_medida=None if (composicao or insumo) and not lido.unidade else unidade(lido.unidade),
            quantidade=Decimal(lido.quantidade),
            preco_unitario=Decimal(lido.preco) if lido.preco is not None else None,
        )
        if item.preco_unitario is None and not (composicao or insumo):
            raise ValueError(f"Linha {lido.linha}: item sem preço e sem código cadastrado.")
        item.save()
    return orcamento


def gerar_modelo(destino):
    """Planilha modelo com o layout reconhecido pela importação."""
    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Orçamento"
    aba.append(["Orçamento - Nome da obra"])
    aba.append(["Linhas sem quantidade são etapas; linhas com quantidade são itens."])
    aba.append([])
    cabecalho = ["Item", "Código", "Descrição", "Und", "Quant.", "Preço unitário", "Total"]
    aba.append(cabecalho)
    linhas = [
        ("1", "", "SERVIÇOS PRELIMINARES", None, None, None),
        ("1.1", "", "Instalação do canteiro de obras", "m2", 120, 85.40),
        ("1.2", "", "Locação da obra", "m2", 450, 6.35),
        ("2", "", "INFRAESTRUTURA", None, None, None),
        ("2.1", "", "Fundações", None, None, None),
        ("2.1.1", "", "Escavação manual de valas", "m3", 38.5, 72.10),
        ("2.1.2", "C-001", "Concreto fck 25 MPa", "m3", 22, 571.55),
        ("2.1.3", "I-007", "Aço CA-50", "kg", 1800, 7.80),
        ("3", "", "ALVENARIA", None, None, None),
        ("3.1", "", "Alvenaria de vedação com bloco cerâmico 9 cm", "m2", 950, 72.14),
    ]
    for numero, (item, codigo, descricao, und, quant, preco) in enumerate(linhas, start=5):
        total = f"=E{numero}*F{numero}" if quant is not None else None
        aba.append([item, codigo, descricao, und, quant, preco, total])
    larguras = {"A": 9, "B": 10, "C": 52, "D": 7, "E": 11, "F": 15, "G": 15}
    for coluna, largura in larguras.items():
        aba.column_dimensions[coluna].width = largura
    negrito = openpyxl.styles.Font(bold=True)
    for celula in aba[4]:
        celula.font = negrito
    aba["A1"].font = openpyxl.styles.Font(bold=True, size=13)
    for numero, (_, _, _, _, quant, _) in enumerate(linhas, start=5):
        if quant is None:
            for celula in aba[numero]:
                celula.font = negrito
        aba.cell(numero, 1).number_format = "@"
        for coluna in (6, 7):
            aba.cell(numero, coluna).number_format = "#,##0.00"
    livro.save(destino)
